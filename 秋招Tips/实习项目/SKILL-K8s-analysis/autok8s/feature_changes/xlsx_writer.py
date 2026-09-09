"""xlsx_writer.py — 写出 "特性变更" sheet, 14 列, 格式对齐参考文件。

列布局 (14 列):
  A 版本变更阶段 | B 变更类型 | C 特性名称 | D 特性阶段变化 | E 默认值变化 |
  F 默认值锁定 | G 是否兼容 | H 兼容分析 | I 分析结论 | J 排查方法 |
  K 参考资料 | L 详细说明 | M 建议开启？ | N 补充说明

样式:
  - 字体: 微软雅黑 11
  - 表头: 加粗白字, theme=7 实色填充; B/C 列左对齐, 其余列居中, 垂直居中, 自动换行
  - 数据: 不加粗; A-L/N 列左对齐, M 列居中; 垂直居中; 自动换行; 白色实色填充
  - 全表细边框; 行高按内容估算; 冻结首行
"""
from __future__ import annotations

from pathlib import Path
import json
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .analyzer import FeatureChange
from .content_store import COMPAT_VALS, ORIG, lookup, suggest_enable
from .research import verify_evidence
from .narrative import finalise_narrative


HEADER = [
    "版本变更阶段", "变更类型", "特性名称", "特性阶段变化", "默认值变化",
    "默认值锁定", "是否兼容", "兼容分析", "分析结论", "排查方法",
    "参考资料", "详细说明", "建议开启？", "补充说明",
]

COLUMN_WIDTHS = [
    50.0, 13.33, 56.0, 19.44, 13.0,
    12.0, 10.22, 22.55, 10.0, 57.89,
    23.55, 50.0, 11.44, 32.0,
]

_FONT_NAME = "微软雅黑"
_FONT_SIZE = 11
_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill(patternType="solid", fgColor=Color(theme=7))
_HEADER_FONT = Font(name=_FONT_NAME, size=_FONT_SIZE, bold=True, color="FFFFFFFF")
_DATA_FILL = PatternFill(patternType="solid", fgColor="FFFFFFFF")
_DATA_FONT = Font(name=_FONT_NAME, size=_FONT_SIZE, bold=False, color="FF000000")

_HEADER_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_HEADER_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_DATA_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_DATA_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

_HEADER_ALIGN = {2: _HEADER_LEFT, 3: _HEADER_LEFT}
_DATA_CENTER_COLS = {13}


def _range_end_stage(stage_change: str) -> str:
    """Return the stage at the end of the requested analysis range."""
    before, separator, after = (stage_change or "").partition("->")
    return (after or before).strip()


def _range_end_default(default_change: str) -> bool:
    """Return the default value at the end of the requested analysis range."""
    _, separator, after = (default_change or "").partition("->")
    return after.strip().lower() == "true" if separator else False


def _estimate_row_height(values: list, widths: list[float]) -> float:
    max_lines = 1
    for val, width in zip(values, widths):
        if val is None:
            continue
        text = str(val)
        if not text:
            continue
        effective = sum(2 if ord(ch) > 127 else 1 for ch in text)
        chars_per_line = max(int(width * 0.9), 4)
        explicit = text.count("\n") + 1
        wrapped = max(1, -(-effective // chars_per_line))
        max_lines = max(max_lines, explicit, wrapped)
    height = max_lines * 15 + 6
    return min(max(height, 24.0), 240.0)


def _fallback_check_method(row, last_stage: str) -> str:
    """无 KEP/注释来源时不兼容特性的最小排查指引 (而非留空)。"""
    name = row.name
    if row.change_type == "Deprecated":
        return f"检查集群是否使用已弃用的 {name} 特性门控，若使用需规划迁移或关闭。"
    if row.change_type == "Added":
        return f"确认集群是否需要 {name} 新特性（{last_stage} 阶段），评估默认值变化的影响。"
    return f"检查集群对 {name} 特性门控的依赖，评估阶段/默认值变化的影响。"


_SECURITY_GROUP_NAMES = frozenset({
    "AuthorizePodWebsocketUpgradeCreatePermission",
    "ConstrainedImpersonation",
    "KubeletFineGrainedAuthz",
    "KubeletEnsureSecretPulledImages",
    "MutatingAdmissionPolicy",
    "PodCertificateRequest",
    "ProcMountType",
    "SELinuxChangePolicy",
    "SELinuxMountReadWriteOncePod",
    "SupplementalGroupsPolicy",
    "UserNamespacesSupport",
    "UserNamespacesHostNetworkSupport",
    "ExtendWebSocketsToKubelet",
    "ExternalServiceAccountTokenSigner",
    "CSIServiceAccountTokenSecrets",
    "StructuredAuthenticationConfigurationJWKSMetrics",
    "ManifestBasedAdmissionControlConfig",
    "RelaxedServiceNameValidation",
    "StrictIPCIDRValidation",
    "ServiceCIDRStatusFieldWiping",
})

_PREFIX_GROUPS = (
    ("DRA", "DRA"),
    ("DynamicResourceAllocation", "DRA"),
    ("Kubelet", "Kubelet"),
    ("StaleControllerConsistency", "StaleControllerConsistency"),
    ("Component", "Component"),
    ("InPlacePod", "InPlacePod"),
    ("Image", "Image"),
    ("DeclarativeValidation", "DeclarativeValidation"),
    ("Mutable", "Mutable"),
    ("Pod", "Pod"),
    ("Disable", "Disable"),
    ("Volume", "Volume"),
    ("Workload", "Workload"),
)

_GROUP_ORDER = {
    "安全/鉴权/准入/权限": 0,
    "DRA": 1,
    "Kubelet": 2,
    "StaleControllerConsistency": 3,
    "Component": 4,
    "InPlacePod": 5,
    "Image": 6,
    "DeclarativeValidation": 7,
    "Mutable": 8,
    "Pod": 9,
    "Disable": 10,
    "Volume": 11,
    "Workload": 12,
    "其他": 13,
}


def _max_minor_version(version_changes: str) -> int:
    minors = re.findall(r"v1\.(\d+)", version_changes or "")
    return max((int(m) for m in minors), default=0)


def _group_sort_key(row: "FeatureChange") -> tuple:
    if row.name in _SECURITY_GROUP_NAMES:
        label = "安全/鉴权/准入/权限"
    else:
        label = "其他"
        for prefix, grp in _PREFIX_GROUPS:
            if row.name.startswith(prefix):
                label = grp
                break
    return (_GROUP_ORDER[label], -_max_minor_version(row.version_changes), row.name)


def _group_and_sort_rows(rows: list) -> list:
    return sorted(rows, key=_group_sort_key)


def _prefetch_research(
    rows: list[FeatureChange], gates: dict, *, version: str = "",
    data_dir: str | Path | None = None, offline: bool = False,
) -> dict[str, dict]:
    from concurrent.futures import ThreadPoolExecutor, as_completed

    tasks: dict[str, tuple] = {}
    for row in rows:
        if row.compat_analysis in COMPAT_VALS:
            continue
        if row.name in ORIG:
            continue
        gate = gates.get(row.name)
        gate_kep = gate.kep if gate else None
        gate_desc = gate.desc if gate else None
        tasks[row.name] = (row.name, gate_kep, gate_desc)

    results: dict[str, dict] = {}
    if not tasks:
        return results

    max_workers = min(8, len(tasks))
    with ThreadPoolExecutor(max_workers=max_workers) as pool:
        futures = {
            pool.submit(
                lookup, name, kep, desc, version=version,
                data_dir=data_dir, offline=offline,
            ): name
            for name, kep, desc in tasks.values()
        }
        for fut in as_completed(futures):
            name = futures[fut]
            try:
                results[name] = fut.result()
            except Exception as e:
                from autok8s.common.logging import get_logger
                get_logger("xlsx_writer").warning("lookup %s 失败: %s", name, e)
                results[name] = {
                    "排查方法": None, "参考资料": None, "详细说明": None,
                    "建议开启": None, "补充说明": None,
                }
    return results


def _enforce_evidence_binding(research: dict, gate) -> None:
    """Ensure 参考资料 (K) exists whenever 详细说明 (L) exists.

    Rule:
    - L requires K (detailed explanation needs a source).
    - J (排查方法) may exist without K when derived from feature-gate
      metadata (name/stage/default), because it is an operational check
      step, not a sourced claim.
    - If K is missing, try gate.kep; if still missing, clear L only.
    - Clear L when it merely restates J.
    Mutates *research* in place.
    """
    j = research.get("排查方法")
    l = research.get("详细说明")
    k = research.get("参考资料")

    if not l:
        return

    if not k:
        if gate and gate.kep:
            research["参考资料"] = gate.kep
        else:
            research["详细说明"] = None
            return

    if l and j:
        norm_l = l.replace("\n", " ").strip()
        norm_j = j.replace("\n", " ").strip()
        if norm_l == norm_j or norm_l in norm_j or norm_j in norm_l:
            research["详细说明"] = None


def _clean_narrative_text(text: str) -> str:
    """Clean machine-translated text: fix known mistranslations, zero-width
    chars, mid-sentence line breaks, stage-prefix leaks, KEP README subsection
    label leaks, and double spaces."""
    if not text:
        return text
    import re as _re
    text = text.replace("\u200b", "").replace("\u200c", "").replace("\u200d", "").replace("\ufeff", "")
    text = text.replace("租赁", "Lease").replace("吊舱", "Pod").replace("豆荚", "Pod")
    text = text.replace("观看流", "watch 流")
    text = text.replace("取决于", "依赖")
    text = _re.sub(r"^(beta|alpha|ga)\s*[：:]\s*v1\.\d+\s*", "", text, flags=_re.IGNORECASE)
    text = text.replace("的功能门", "特性门").replace("此功能门", "此特性门").replace("功能门以恢复", "特性门以恢复")
    text = text.replace("功能门下", "特性门下").replace("功能门", "特性门")
    text = _re.sub(r"用户故事（可选）\s*", "", text)
    text = _re.sub(r"用户故事\s*", "", text)
    text = _re.sub(r"故事[一二三1-9]?\s*", "", text)
    text = _re.sub(r"简化的资源管理\s*", "", text)
    text = text.replace("\n\n", "\x00PARA\x00")
    text = text.replace("\n", " ")
    text = text.replace("\x00PARA\x00", "\n\n")
    text = _re.sub(r" {2,}", " ", text)
    return text.strip()


def _strip_template_detail(research: dict, row: FeatureChange) -> None:
    """Remove low-quality 详细说明 that merely restates D/E columns.

    The fallback narrative produces text like "X 的已解析变更：阶段升级为…"
    which adds no information beyond the mechanical columns.  Clear it so
    the column stays empty unless real sourced content is available.
    """
    l = research.get("详细说明")
    if not l:
        return
    if "已解析变更" in l:
        research["详细说明"] = None


def _fill_supplement(research: dict, row: FeatureChange, gate, *, offline: bool = False) -> None:
    """Populate 补充说明 (N) when 详细说明 (L) is missing, for non-compatible rows only.

    N carries per-feature functional context so each row is distinct:
    1. If the go file has a description for this gate, translate it to Chinese.
    2. If not, synthesize one sentence from the feature name and change type.
    """
    if research.get("补充说明"):
        return
    if row.compat_analysis in COMPAT_VALS:
        return
    l = research.get("详细说明")
    if l:
        return

    desc_zh = None

    if gate and gate.desc and not offline:
        from .fetcher import translate_to_chinese as _translate
        import re as _re
        raw_desc = " ".join(gate.desc).strip()
        if raw_desc and len(raw_desc) > 10:
            cn_desc, _ = _translate(raw_desc[:400])
            if cn_desc:
                cn_desc = _clean_narrative_text(cn_desc)
                if cn_desc:
                    desc_zh = cn_desc

    if not desc_zh:
        desc_zh = f"{row.name} 特性门控发生 {row.change_type} 变更，具体功能以参考资料为准。"

    research["补充说明"] = desc_zh[:300]


def write_workbook(
    rows: list[FeatureChange],
    gates: dict,
    out_path: str | Path,
    sheet_name: str = "特性变更",
    *,
    version: str = "",
    data_dir: str | Path | None = None,
    offline: bool = False,
    audit_path: str | Path | None = None,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(HEADER)

    ordered_rows = _group_and_sort_rows(rows)
    prefetch = _prefetch_research(ordered_rows, gates, version=version, data_dir=data_dir, offline=offline)
    all_research: dict[str, dict] = dict(prefetch)

    for row in ordered_rows:
        gate = gates.get(row.name)
        gate_kep = gate.kep if gate else None
        gate_desc = gate.desc if gate else None
        if row.compat_analysis in COMPAT_VALS:
            research = {
                "排查方法": None, "参考资料": None, "详细说明": None,
                "建议开启": None, "补充说明": None,
            }
        else:
            research = prefetch.get(row.name) or lookup(
                row.name, gate_kep, gate_desc, version=version,
                data_dir=data_dir, offline=offline,
            )
        all_research[row.name] = research
        # Use the analyzed range result.  gate.specs may include future
        # versions and must not leak a later stage into this workbook.
        last_stage = _range_end_stage(row.stage_change)
        last_default = _range_end_default(row.default_change)
        if not row.compat_analysis:
            final = finalise_narrative(
                row,
                last_stage,
                research.get("排查方法"),
                research.get("详细说明"),
                research.get("补充说明"),
            )
            research.update(final)
        # 兼容分析有内容 (无影响行) 时, 其后各列 (分析结论/排查方法/参考资料/
        # 详细说明/建议开启/补充说明) 必须全部留空, 因此不再生成建议开启值。
        if research.get("建议开启") is None and not row.compat_analysis:
            research["建议开启"] = suggest_enable(row.change_type, last_stage, last_default)

        for _key in ("排查方法", "详细说明", "补充说明"):
            if research.get(_key):
                research[_key] = _clean_narrative_text(research[_key])
        _strip_template_detail(research, row)
        _enforce_evidence_binding(research, gate)
        _fill_supplement(research, row, gate, offline=offline)

        ws.append([
            row.version_changes,
            row.change_type,
            row.name,
            row.stage_change,
            row.default_change,
            row.lock_change,
            row.compatible,
            row.compat_analysis,
            None,
            research["排查方法"],
            research["参考资料"],
            research["详细说明"],
            research["建议开启"],
            research["补充说明"],
        ])

    for idx, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    for r in range(1, ws.max_row + 1):
        is_header = r == 1
        for c in range(1, len(HEADER) + 1):
            cell = ws.cell(r, c)
            cell.border = _BORDER
            if is_header:
                cell.alignment = _HEADER_ALIGN.get(c, _HEADER_CENTER)
                cell.fill = _HEADER_FILL
                cell.font = _HEADER_FONT
            else:
                cell.alignment = _DATA_CENTER if c in _DATA_CENTER_COLS else _DATA_LEFT
                cell.fill = _DATA_FILL
                cell.font = _DATA_FONT
        if is_header:
            ws.row_dimensions[r].height = 25.2
        else:
            values = [ws.cell(r, c).value for c in range(1, len(HEADER) + 1)]
            ws.row_dimensions[r].height = _estimate_row_height(values, COLUMN_WIDTHS)

    ws.freeze_panes = "A2"

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    if audit_path is not None:
        audit = Path(audit_path)
        audit.parent.mkdir(parents=True, exist_ok=True)
        audit.write_text(
            json.dumps(verify_evidence(ordered_rows, all_research), ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
