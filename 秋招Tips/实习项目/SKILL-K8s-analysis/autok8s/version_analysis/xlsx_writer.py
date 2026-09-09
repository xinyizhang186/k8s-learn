"""xlsx_writer.py — 生成版本分析 xlsx, 格式与参考文件一致。

单 sheet "版本分析", 含两个部分:
  1. 关键特性价值分析 (标题行 + 列头 + 特性数据行)
  2. 关键变更风险分析 (标题行 + 列头 + 弃用数据行)

样式:
  - 字体: 微软雅黑
  - 标题: 16 bold; 段标题: 14 bold; 列头: 12 bold; 数据: 12
  - "现状：" 和 "本特性增强：" 使用富文本加粗
  - 全表细边框, 自动换行, 垂直居中
  - 数据行不设固定行高, 由 Excel 按 wrap_text 内容自动撑开 (文本完全显示)
  - 列宽: [22.0, 55.0, 120.0, 18.0, 68.0]
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Side, Font
from openpyxl.utils import get_column_letter


_FONT = "微软雅黑"
_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

COLUMN_WIDTHS = [22.0, 55.0, 120.0, 18.0, 68.0]

_BOLD = InlineFont(rFont=_FONT, sz=12, b=True)
_NORMAL = InlineFont(rFont=_FONT, sz=12, b=False)

_F_TITLE = Font(name=_FONT, size=16, bold=True)
_F_SECTION = Font(name=_FONT, size=14, bold=True)
_F_HEADER = Font(name=_FONT, size=12, bold=True)
_F_DATA = Font(name=_FONT, size=12, bold=False)

_A_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
_A_WRAP_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

_FEATURE_HEADER = ["分类", "特性名称", "特性功能介绍", "特性价值领域", "特性功能价值分析"]
_DEPRECATION_HEADER = ["分类", "变更名称", "风险详细描述", "风险涉及领域", "技术or商业影响"]


def _make_rich_text(text: str) -> CellRichText:
    """将含 '现状：' 和 '本特性增强：' 的文本转为富文本, 标签加粗。"""
    blocks = []
    remaining = text

    for label in ["现状：", "本特性增强："]:
        idx = remaining.find(label)
        if idx >= 0:
            if idx > 0:
                blocks.append(TextBlock(_NORMAL, remaining[:idx]))
            blocks.append(TextBlock(_BOLD, label))
            remaining = remaining[idx + len(label):]

    if remaining:
        blocks.append(TextBlock(_NORMAL, remaining))

    if not blocks:
        return CellRichText([TextBlock(_NORMAL, text)])
    return CellRichText(blocks)


def write_workbook(analysis: dict, out_path: str | Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "版本分析"

    title = analysis["title"]
    features = analysis["features"]
    deprecations = analysis["deprecations"]

    col_count = 5

    row = 1

    ws.cell(row, 1, title)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
    _style_cell(ws.cell(row, 1), _F_TITLE, _A_WRAP_CENTER)
    ws.row_dimensions[row].height = 23.4
    row += 1

    ws.cell(row, 1, "关键特性价值分析")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
    _style_cell(ws.cell(row, 1), _F_SECTION, _A_WRAP_CENTER)
    ws.row_dimensions[row].height = 33.0
    row += 1

    for c, h in enumerate(_FEATURE_HEADER, 1):
        ws.cell(row, c, h)
        _style_cell(ws.cell(row, c), _F_HEADER, _A_WRAP_CENTER)
    ws.row_dimensions[row].height = 34.5
    row += 1

    for feat in features:
        vals = [feat.get(k, "") for k in _FEATURE_HEADER]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row, c)
            if c == 3 and isinstance(v, str) and ("现状：" in v or "本特性增强：" in v):
                cell.value = _make_rich_text(v)
            else:
                cell.value = v
            _style_cell(cell, _F_DATA, _A_WRAP)
        row += 1

    row += 2

    ws.cell(row, 1, "关键变更风险分析")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
    _style_cell(ws.cell(row, 1), _F_SECTION, _A_WRAP_CENTER)
    ws.row_dimensions[row].height = 33.0
    row += 1

    for c, h in enumerate(_DEPRECATION_HEADER, 1):
        ws.cell(row, c, h)
        _style_cell(ws.cell(row, c), _F_HEADER, _A_WRAP_CENTER)
    ws.row_dimensions[row].height = 30.0
    row += 1

    for dep in deprecations:
        vals = [dep.get(k, "") for k in _DEPRECATION_HEADER]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row, c, v)
            _style_cell(cell, _F_DATA, _A_WRAP)
        row += 1

    for c, w in enumerate(COLUMN_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(c)].width = w

    for r in range(1, row):
        for c in range(1, col_count + 1):
            ws.cell(r, c).border = _BORDER

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def _style_cell(cell, font, alignment):
    cell.font = font
    cell.alignment = alignment
    cell.border = _BORDER
