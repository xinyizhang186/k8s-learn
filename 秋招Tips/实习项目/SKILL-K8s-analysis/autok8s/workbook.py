"""workbook.py — 双 Sheet 合并工作簿生成 (版本分析 + 特性变更)。

将 version_analysis 和 feature_changes 两个分析结果合并到单个 xlsx,
含两个工作表: "版本分析" 和 "特性变更"。
输出前校验两个 Sheet 名和列头, 确保符合人工审查模板。
"""
from __future__ import annotations

import tempfile
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText
from openpyxl.cell.cell import MergedCell

from .version_analysis.content_gen import generate_analysis, generate_multi_version_analysis
from .version_analysis.fetcher import fetch_release_blog
from .version_analysis.xlsx_writer import write_workbook
from .feature_changes.generator import generate_workbook as generate_feature_changes


_EXPECTED_VERSION_HEADER = ["分类", "特性名称", "特性功能介绍", "特性价值领域", "特性功能价值分析"]
_EXPECTED_CHANGE_HEADER = [
    "版本变更阶段", "变更类型", "特性名称", "特性阶段变化", "默认值变化",
    "默认值锁定", "是否兼容", "兼容分析", "分析结论", "排查方法",
    "参考资料", "详细说明", "建议开启？", "补充说明",
]


def _version_tuple(value: str) -> tuple[int, int]:
    parts = value.strip().lstrip("v").split(".")
    if len(parts) < 2:
        raise ValueError(f"版本号格式错误, 应为 '1.34' 形式: {value!r}")
    return int(parts[0]), int(parts[1])


def _copy_worksheet(source, target) -> None:
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    target.freeze_panes = source.freeze_panes
    for merged in source.merged_cells.ranges:
        target.merge_cells(str(merged))
    for key, dimension in source.column_dimensions.items():
        target.column_dimensions[key] = copy(dimension)
    for index, dimension in source.row_dimensions.items():
        target.row_dimensions[index] = copy(dimension)
    for row in source.iter_rows():
        for cell in row:
            copied = target[cell.coordinate]
            if isinstance(copied, MergedCell):
                continue
            if isinstance(cell.value, CellRichText):
                copied.value = cell.value
            else:
                copied.value = copy(cell.value)
            if cell.has_style:
                copied._style = copy(cell._style)
            if cell.number_format:
                copied.number_format = cell.number_format
            copied.alignment = copy(cell.alignment)
            copied.protection = copy(cell.protection)
            copied.font = copy(cell.font)
            copied.fill = copy(cell.fill)
            copied.border = copy(cell.border)
    target.sheet_properties = copy(source.sheet_properties)
    target.sheet_format = copy(source.sheet_format)
    target.page_margins = copy(source.page_margins)
    target.page_setup = copy(source.page_setup)


def _assert_contract(workbook) -> None:
    if workbook.sheetnames != ["版本分析", "特性变更"]:
        raise ValueError(f"工作表不符合模板约定: {workbook.sheetnames}")
    version_header = [cell.value for cell in workbook["版本分析"][3]]
    change_header = [cell.value for cell in workbook["特性变更"][1]]
    if version_header != _EXPECTED_VERSION_HEADER:
        raise ValueError(f"版本分析列不符合模板约定: {version_header}")
    if change_header != _EXPECTED_CHANGE_HEADER:
        raise ValueError(f"特性变更列不符合模板约定: {change_header}")


def generate_combined_workbook(
    blog_urls: list[str],
    go_file: str | Path,
    lo: tuple[int, int],
    hi: tuple[int, int],
    output: str | Path,
    *,
    data_dir: str | Path | None = None,
    offline: bool = False,
    audit_path: str | Path | None = None,
) -> tuple[int, int]:
    blogs = []
    for url in blog_urls:
        blog = fetch_release_blog(url)
        if not blog:
            raise RuntimeError(f"无法抓取或解析发布博客: {url}")
        blogs.append(blog)
    blogs.sort(key=lambda item: _version_tuple(item.version), reverse=True)
    analysis = generate_analysis(blogs[0]) if len(blogs) == 1 else generate_multi_version_analysis(blogs)

    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="autok8s-") as directory:
        tmp = Path(directory)
        version_path = tmp / "version.xlsx"
        change_path = tmp / "changes.xlsx"
        write_workbook(analysis, version_path)
        change_count = generate_feature_changes(
            go_file, change_path, lo=lo, hi=hi, data_dir=data_dir,
            offline=offline, audit_path=audit_path,
        )

        destination = load_workbook(version_path, rich_text=True)
        source = load_workbook(change_path, rich_text=True)
        change_sheet = destination.create_sheet("特性变更")
        _copy_worksheet(source.active, change_sheet)
        _assert_contract(destination)
        destination.save(output)

    feature_count = len(analysis["features"])
    return feature_count, change_count
